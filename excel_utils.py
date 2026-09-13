import sqlite3
from openpyxl import Workbook, load_workbook
from datetime import datetime


def exportar_a_excel(ruta=None):
    """Exporta todas las citas a un archivo Excel."""
    if ruta is None:
        ruta = f"citas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    conexion = sqlite3.connect("agenda.db")
    cursor = conexion.cursor()
    cursor.execute("SELECT id, cliente, telefono, fecha, hora, motivo, estado FROM citas ORDER BY fecha, hora")
    filas = cursor.fetchall()
    conexion.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Citas"

    encabezados = ["ID", "Cliente", "Telefono", "Fecha", "Hora", "Motivo", "Estado"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = celda.font.copy(bold=True)

    for fila in filas:
        ws.append(fila)

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_largo = max((len(str(c.value)) if c.value else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_largo + 4

    wb.save(ruta)
    return ruta


def importar_desde_excel(ruta):
    """Importa citas desde un archivo Excel. Devuelve (importadas, errores)."""
    wb = load_workbook(ruta)
    ws = wb.active

    conexion = sqlite3.connect("agenda.db")
    cursor = conexion.cursor()

    importadas = 0
    errores = 0

    # Saltar encabezado (fila 1)
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not fila or not fila[1]:
            continue
        try:
            datos = (list(fila) + [None] * 7)[:7]
            _, cliente, telefono, fecha, hora, motivo, estado = datos

            fecha = str(fecha).strip() if fecha else ""
            hora = str(hora).strip() if hora else ""

            if not cliente or not fecha or not hora:
                errores += 1
                continue

            cursor.execute(
                "INSERT INTO citas (cliente, telefono, fecha, hora, motivo, estado) VALUES (?, ?, ?, ?, ?, ?)",
                (cliente, telefono or "", fecha, hora, motivo or "", estado or "pendiente")
            )
            importadas += 1
        except Exception:
            errores += 1

    conexion.commit()
    conexion.close()
    return importadas, errores
